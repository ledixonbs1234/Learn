# Coder/.skills/excel-handler/scripts/edit_excel.py
import sys
import re
import subprocess
import importlib
from pathlib import Path

# ĐẢM BẢO HOÀN TOÀN ĐẦU RA LÀ UTF-8
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

def ensure_dependencies():
    required = ["pandas", "openpyxl", "xlrd", "xlwt", "xlutils"]
    missing = []
    for pkg in required:
        try:
            importlib.import_module(pkg)
        except ImportError:
            missing.append(pkg)
            
    if not missing:
        return

    installed = False

    # 1. Tầng 1: Sử dụng uv pip cài đặt trực tiếp vào .venv hiện tại
    try:
        res = subprocess.run(
            ["uv", "pip", "install", "--python", sys.executable] + missing + ["--quiet"],
            capture_output=True,
            text=True
        )
        if res.returncode == 0:
            installed = True
    except FileNotFoundError:
        pass
    except Exception:
        pass

    # 2. Tầng 2: Dự phòng sử dụng pip tiêu chuẩn của môi trường ảo
    if not installed:
        try:
            res = subprocess.run(
                [sys.executable, "-m", "pip", "install"] + missing + ["--quiet"],
                capture_output=True,
                text=True
            )
            if res.returncode == 0:
                installed = True
        except Exception:
            pass

    # 3. Tầng 3: Tái sinh pip bằng ensurepip
    if not installed:
        try:
            import ensurepip
            ensurepip.bootstrap(upgrade=True, default_pip=True)
            res = subprocess.run(
                [sys.executable, "-m", "pip", "install"] + missing + ["--quiet"],
                capture_output=True,
                text=True
            )
            if res.returncode == 0:
                installed = True
        except Exception:
            pass

    # 4. Tầng 4: Dự phòng trình thông dịch bên ngoài
    if not installed:
        for host_python in ["py", "python", "python3"]:
            try:
                subprocess.run(
                    [host_python, "-m", "pip", "install"] + missing + ["--quiet"],
                    capture_output=True,
                    check=True
                )
                all_imported = True
                for pkg in missing:
                    try:
                        importlib.import_module(pkg)
                    except ImportError:
                        all_imported = False
                if all_imported:
                    return
            except Exception:
                pass

ensure_dependencies()

def parse_coordinate(coord_str: str):
    match = re.match(r"^([a-zA-Z]+)([0-9]+)$", coord_str)
    if not match:
        raise ValueError(f"Invalid coordinate format '{coord_str}'")
    col_str, row_str = match.groups()
    col_idx = 0
    for char in col_str.upper():
        col_idx = col_idx * 26 + (ord(char) - ord('A') + 1)
    col_idx -= 1
    row_idx = int(row_str) - 1
    return row_idx, col_idx

def edit_excel(file_path: str, sheet_name: str, coordinate: str, value: str):
    try:
        path = Path(file_path).resolve()
        if not path.exists():
            print(f"Error: File not found at {file_path}")
            sys.exit(1)
            
        ext = path.suffix.lower()
        if value.replace('.', '', 1).isdigit():
            val = float(value) if '.' in value else int(value)
        else:
            val = value
            
        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path)
            if sheet_name not in wb.sheetnames:
                print(f"Error: Sheet '{sheet_name}' not found.")
                sys.exit(1)
            ws = wb[sheet_name]
            ws[coordinate] = val
            wb.save(path)
            wb.close()
            print(f"Successfully edited cell {coordinate} on sheet '{sheet_name}' with value: '{value}' (.xlsx)")
            
        elif ext == ".xls":
            import xlrd
            from xlutils.copy import copy
            
            row_idx, col_idx = parse_coordinate(coordinate)
            rb = xlrd.open_workbook(path, formatting_info=True)
            if sheet_name not in rb.sheet_names():
                print(f"Error: Sheet '{sheet_name}' not found.")
                sys.exit(1)
                
            sheet_idx = rb.sheet_names().index(sheet_name)
            wb = copy(rb)
            ws = wb.get_sheet(sheet_idx)
            ws.write(row_idx, col_idx, val)
            wb.save(path)
            print(f"Successfully edited cell {coordinate} on sheet '{sheet_name}' with value: '{value}' (.xls)")
        else:
            print(f"Error: Unsupported file extension '{ext}'")
            sys.exit(1)
            
    except Exception as e:
        print(f"Error editing cell: {str(e)}")
        sys.exit(1)

if __name__ == "__main__":
    if len(sys.argv) < 5:
        print("Usage: python edit_excel.py <file_path> <sheet_name> <coordinate> <value>")
        sys.exit(1)
    edit_excel(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4])